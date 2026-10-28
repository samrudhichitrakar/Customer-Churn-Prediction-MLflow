from pathlib import Path
import json,sys,subprocess
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
R=Path(__file__).resolve().parent; f=R/"outputs"/"results.json"
if not f.exists(): print("Run python run_project.py first."); sys.exit(1)
d=json.loads(f.read_text()); m=d["metrics"]
doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
doc.styles["Normal"].font.name="Arial"; doc.styles["Normal"].font.size=Pt(10)
p=doc.add_heading("MLflow Installation, Experiment Tracking and Data Drift Analysis",0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for t in ["ISE-2 Practical Report | Academic Year 2026–27","Project: Customer Churn Prediction using the IBM Telco Customer Churn Dataset"]:
 p=doc.add_paragraph(t); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_heading("Student Details",1)
for x in ["Student Name: Pranali Shashikant Wadtile","Roll Number: DS08","Department / Class: CSE (Data Science), Fourth Year","Academic Year: 2026–27","GitHub Repository URL: [Paste your actual repository URL after creating it]"]: doc.add_paragraph(x)
doc.add_heading("1. MLflow Installation and Experiment Tracking",1); doc.add_heading("1.1 Environment Setup",2)
for x in ["Operating System: Windows 10/11 (confirm exact version).","Python Version: record using python --version.","MLflow Version: record using mlflow --version.","Install dependencies with pip install -r requirements.txt. Run python run_project.py and confirm the run is saved in the local mlruns directory and appears in the MLflow UI."]: doc.add_paragraph(x,style="List Bullet")
doc.add_heading("1.2 Experiment Tracking",2)
for x in ["Experiment Name: ISE2_Customer_Churn_Prediction.","Objective: Predict telecom customer churn and track model parameters, metrics, model artifact, and drift results.","Model: RandomForestClassifier; 150 trees; maximum depth 12; class_weight=balanced; random_state=42.",
f"Test metrics: Accuracy={m['accuracy']:.4f}; Precision={m['precision']:.4f}; Recall={m['recall']:.4f}; F1={m['f1']:.4f}; ROC-AUC={m['roc_auc']:.4f}.",
f"Rows={m['total_rows']}; training rows={m['training_rows']}; test rows={m['test_rows']}; input features={m['feature_count']}.",
"Run name: RandomForest_with_simulated_churn_drift. The model and drift_results.csv are logged as MLflow artifacts."]: doc.add_paragraph(x,style="List Bullet")
doc.add_heading("2. Dataset Identification",1); doc.add_heading("2.1 Problem Statement",2); doc.add_paragraph(d["problem"])
doc.add_heading("2.2 Dataset Details",2)
for x in [f"Dataset: {d['dataset_name']}.",f"Source CSV: {d['dataset_source']}.",f"Number of records: {m['total_rows']}.",f"Number of input features: {m['feature_count']}.","Target: Churn (Yes=1, No=0).","Example features: tenure, MonthlyCharges, TotalCharges, Contract, InternetService, PaymentMethod and service subscription indicators.","Use case: Supervised learning / binary classification."]: doc.add_paragraph(x,style="List Bullet")
doc.add_heading("2.3 Dataset Selection Rationale",2); doc.add_paragraph("Customer churn is a practical business problem. Predicting customers who may leave can help a telecom provider investigate service issues and plan retention activities. The dataset contains both numeric and categorical attributes, making it useful for demonstrating preprocessing and data monitoring.")
doc.add_heading("3. Drift Scenario and Detection Method",1); doc.add_heading("3.1 Drift Type",2); doc.add_paragraph("Simulated data drift: changes in input feature distributions. This analysis does not establish concept drift.")
doc.add_heading("3.2 Reference and Current Data",2); doc.add_paragraph("The data is split into training and test sets using a stratified 75/25 split and random_state=42. Test features form the reference data. A copy is modified to simulate current data: MonthlyCharges is increased by 15%, tenure is reduced by six months (minimum zero), and Contract categories are shifted. These changes are intentional and are not observed production drift.")
doc.add_heading("3.3 Detection Method",2)
for x in ["Numeric features: two-sample Kolmogorov–Smirnov (KS) test.","Categorical features: chi-square test of category counts.","Threshold: p-value < 0.05 flags a feature.","Per-feature results are saved in outputs/drift_results.csv and visualized in outputs/drift_plot.png. Multiple tests make this an educational demonstration rather than a complete production monitoring policy."]: doc.add_paragraph(x,style="List Bullet")
doc.add_heading("4. Results, Interpretation and Conclusion",1); doc.add_heading("4.1 Drift Analysis Results",2)
doc.add_paragraph(f"Dataset churn rate={m['churn_rate']:.4f}. Test metrics: Accuracy={m['accuracy']:.4f}, Precision={m['precision']:.4f}, Recall={m['recall']:.4f}, F1={m['f1']:.4f}, ROC-AUC={m['roc_auc']:.4f}. Drift was detected in {m['drifted_feature_count']} of {m['feature_count']} input features: {', '.join(m['drifted_features']) if m['drifted_features'] else 'none'}. See outputs/drift_results.csv and outputs/drift_plot.png.")
doc.add_heading("4.2 Interpretation",2); doc.add_paragraph("Accuracy is the overall proportion of correct predictions. Precision measures how many predicted churners actually churned; recall measures how many actual churners were identified; F1 balances precision and recall; ROC-AUC summarizes ranking performance across thresholds. A drift flag means the feature distribution changed; it does not by itself prove model performance has degraded.")
doc.add_heading("4.3 Conclusion and Recommended Action",2); doc.add_paragraph("This project demonstrates MLflow tracking for a customer churn classifier and statistical tests for simulated input drift. In a real telecom environment, investigate data-pipeline or customer-mix changes, validate data quality, and assess performance on recent labeled observations before deciding whether to retrain or change retention actions.")
doc.add_heading("5. GitHub Repository and Reproducibility",1)
for x in ["GitHub Repository URL: [Add the URL after creating and uploading to your repository].","Source code: run_project.py; report generator: generate_report.py.","README: README.md; dependencies: requirements.txt.","Generated artifacts: outputs/results.json, outputs/drift_results.csv, outputs/drift_plot.png, outputs/confusion_matrix.png; MLflow tracking store: mlruns/."]: doc.add_paragraph(x,style="List Bullet")
doc.add_heading("Execution Commands",1)
for x in ["py -m venv .venv",r".\.venv\Scripts\Activate.ps1","pip install -r requirements.txt","python run_project.py","python generate_report.py","mlflow ui --backend-store-uri ./mlruns --host 127.0.0.1 --port 5000"]:
 p=doc.add_paragraph(); r=p.add_run(x); r.font.name="Consolas"; r.font.size=Pt(9)
doc.add_heading("References",1)
for x in ["IBM Telco Customer Churn dataset repository: https://github.com/IBM/telco-customer-churn-on-icp4d","MLflow documentation: https://mlflow.org/docs/latest/","scikit-learn documentation: https://scikit-learn.org/stable/","SciPy statistical tests: https://docs.scipy.org/doc/scipy/reference/stats.html"]: doc.add_paragraph(x)
out=R/"ISE2_Customer_Churn_Report_COMPLETED.docx"; doc.save(out); print(f"Report generated: {out}")
